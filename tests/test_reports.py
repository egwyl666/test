"""Отчёты: выручка и прибыль без отменённых, закупка на дату заказа («≈», если истории нет), валюта, комиссия Prom,
ABC, товары без продаж, пустой период, Excel и страница."""

import io
from datetime import datetime, timedelta, timezone

import openpyxl

from promloader import db, orders, products, reports


def at(days_ago: float) -> str:
    return (datetime.now(timezone.utc) - timedelta(days=days_ago)).isoformat()


def product(name, prom_id, cost=None, currency="UAH", price=100, **extra):
    pid = products.create({"name": name, "price": price, "cost_price": cost, "cost_currency": currency, **extra})
    with db.tx() as c:
        c.execute("UPDATE products SET prom_id = ?, synced_at = ?, status = 'synced' WHERE id = ?", (prom_id, db.now(), pid))
    return pid


def order(oid, days_ago, items, status="delivered", commission="0", cancel=None):
    """items: [(prom_id, qty, total)]. Внешний ID в заказах Prom обычно пустой — как и здесь."""
    total = sum(t for _, _, t in items)
    orders._store([{
        "id": oid, "status": status, "date_created": at(days_ago), "price": f"{total:g} грн", "full_price": f"{total:g} грн",
        "cpa_commission": {"amount": commission, "is_refunded": False},
        "cancellation": {"title": cancel, "initiator": "user"} if cancel else None,
        "products": [{"id": pid, "external_id": None, "name": f"Товар {pid}", "quantity": q, "price": f"{t / q:g} грн",
                      "total_price": f"{t:g} грн"} for pid, q, t in items],
    }], first_run=True)


def period():
    return reports.period_bounds("30d")


def test_sales_exclude_canceled_and_count_profit_with_commission():
    product("Фонарь", 501, cost=60)
    product("Повербанк", 502, cost=200)
    order(1, 2, [(501, 2, 200)], commission="20")             # прибыль 200 − 120 − 20 = 60
    order(2, 3, [(502, 1, 300)], commission="30")             # 300 − 200 − 30 = 70
    order(3, 1, [(501, 5, 500)], status="canceled", cancel="Товар не нужен")
    order(4, 1, [(999, 1, 100)])                              # товара нет в программе — выручка есть, прибыль неизвестна
    r = reports.sales(*period())
    assert r["orders"] == 3 and r["revenue"] == 600 and r["avg_check"] == 200 and r["units"] == 4
    assert r["profit"] == 130 and r["commission"] == 50
    assert r["profit_coverage"] == 83  # 500 из 600 грн выручки — с известной закупкой
    assert r["canceled"] == 1 and r["cancel_rate"] == 25.0
    assert r["cancel_reasons"] == [{"reason": "Товар не нужен", "count": 1}]
    assert len(r["days"]) == 30 and sum(d["revenue"] for d in r["days"]) == 600
    assert r["approx"]  # закупка ни разу не менялась — взята текущая


def test_cost_on_order_date_from_journal_and_currency():
    pid = product("Инвертор", 601, cost=10, currency="USD")  # курс в тестах — 40 грн
    order(1, 5, [(601, 1, 600)])
    products.update(pid, {"cost_price": 12})                  # закупка выросла уже после заказа
    r = reports.sales(*period())
    assert r["profit"] == 600 - 400 and not r["approx"]       # 10 $ × 40, а не новые 12 $
    order(2, 0, [(601, 1, 600)])                              # заказ после изменения — по новой закупке
    top = reports.products_report(*period())["top"][0]
    assert top["profit"] == (600 - 400) + (600 - 480) and top["units"] == 2


def test_abc_and_dead_stock():
    for i in range(1, 6):
        product(f"Товар {i}", 700 + i, cost=1)
    with db.tx() as c:  # «старые» товары — им пора было продаваться
        c.execute("UPDATE products SET created_at = ?", (at(60),))
    order(1, 1, [(701, 1, 800)])
    order(2, 1, [(702, 1, 150)])
    order(3, 1, [(703, 1, 50)])
    order(4, 45, [(704, 1, 10)])                              # давно — значит, без продаж 30 дней
    r = reports.products_report(*period())
    abc = {t["name"]: t["abc"] for t in r["top"]}
    assert abc == {"Товар 1": "A", "Товар 2": "B", "Товар 3": "C"} and r["abc"] == {"A": 1, "B": 1, "C": 1}
    dead = {d["name"]: d["last_sale"] for d in r["dead"]}
    assert set(dead) == {"Товар 4", "Товар 5"} and dead["Товар 5"] == ""


def test_suppliers_groups_store_and_empty_period():
    product("Фонарь", 501, cost=60, group_name="Ліхтарі")
    order(1, 1, [(501, 1, 100)])
    s = reports.suppliers_report(*period())
    assert s["groups"][0]["name"] == "Ліхтарі" and s["groups"][0]["profit"] == 40
    store = reports.store_report(*period())
    assert store["total"] == 1 and store["statuses"] == {"synced": 1} and store["quality"]["no_photo"] == 1
    empty = reports.sales(*reports.period_bounds("custom", "2020-01-01", "2020-01-02"))
    assert empty["orders"] == 0 and empty["profit"] is None and len(empty["days"]) == 2


def test_api_xlsx_and_page(client):
    product("Фонарь", 501, cost=60)
    order(1, 1, [(501, 1, 100)])
    r = client.get("/api/reports/sales", params={"period": "7d"}).json()
    assert r["revenue"] == 100 and r["tab"] == "sales"
    for tab in ("sales", "products", "suppliers", "store"):
        res = client.get(f"/api/reports/{tab}.xlsx", params={"period": "30d"})
        assert res.status_code == 200 and "attachment" in res.headers["content-disposition"]
        assert openpyxl.load_workbook(io.BytesIO(res.content)).sheetnames
    assert client.get("/api/reports/nope").status_code == 400
    assert client.get("/api/reports/sales", params={"period": "custom", "date_from": "x"}).status_code == 400
    assert client.get("/reports").status_code == 200


# ---------- 2.17.0: выручка на странице «Заказы», статистика товаров, сводка за день ----------

def test_order_summary_today_yesterday_week_month():
    product("Фонарь", 501, cost=60)
    order(1, 0, [(501, 1, 100)])
    order(2, 0, [(501, 2, 200)], status="canceled")
    order(3, 1, [(501, 3, 300)])
    s = reports.order_summary()
    assert s["today"]["orders"] == 1 and s["today"]["revenue"] == 100 and s["today"]["canceled"] == 1
    assert s["today"]["profit"] == 40
    assert s["7d"]["revenue"] == 400 and s["7d"]["units"] == 4
    assert s["yesterday"]["revenue"] == 300


def test_order_stats_per_product_sort_search_and_share():
    product("Фонарь", 501, cost=60)
    product("Повербанк", 502, cost=200)
    order(1, 2, [(501, 2, 200), (502, 1, 300)])
    order(2, 3, [(501, 1, 100)])
    order(3, 1, [(502, 1, 300)], status="canceled")
    r = reports.order_stats("30d")
    assert r["totals"]["revenue"] == 600 and r["totals"]["canceled"] == 1
    names = [i["name"] for i in r["items"]]
    assert names == ["Повербанк", "Фонарь"]  # по выручке: 300 = 300 → по имени
    fon = r["items"][1]
    assert fon["units"] == 3 and fon["orders"] == 2 and fon["profit"] == 120 and fon["share"] == 50.0
    pb = r["items"][0]
    assert pb["canceled"] == 1 and pb["units"] == 1
    assert [i["name"] for i in reports.order_stats("30d", sort="units")["items"]][0] == "Фонарь"
    assert [i["name"] for i in reports.order_stats("30d", q="повер")["items"]] == ["Повербанк"]


def test_order_stats_api_and_xlsx(client):
    product("Фонарь", 501, cost=60)
    order(1, 1, [(501, 2, 200)])
    r = client.get("/api/orders/stats?period=7d").json()
    assert r["items"][0]["revenue"] == 200
    assert client.get("/api/orders/summary").json()["7d"]["revenue"] == 200
    x = client.get("/api/orders/stats.xlsx?period=7d")
    ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
    rows = list(ws.values)
    assert rows[1][0] == "Товар" and rows[2][0] == "Фонарь" and rows[2][6] == 200
    assert client.get("/api/orders/stats?period=custom&date_from=bad").status_code == 400


def test_daily_summary_once_a_day_after_hour(monkeypatch):
    from promloader import notify
    product("Фонарь", 501, cost=60)
    order(1, 0, [(501, 2, 200)])
    sent = []
    monkeypatch.setattr(notify, "recipients", lambda: [{"chat_id": "1", "events": ["daily_summary"]}])
    monkeypatch.setattr(notify, "send", lambda event, text, background=True: sent.append((event, text)))
    now = datetime.now()
    assert not reports.maybe_daily_summary(now.replace(hour=reports.SUMMARY_HOUR - 1))
    assert reports.maybe_daily_summary(now.replace(hour=reports.SUMMARY_HOUR))
    assert not reports.maybe_daily_summary(now.replace(hour=reports.SUMMARY_HOUR + 1))  # второй раз — нет
    event, text = sent[0]
    assert event == "daily_summary" and "Заказов: <b>1</b>" in text and "Фонарь" in text and "Прибыль" in text


def test_daily_summary_skipped_without_subscribers(monkeypatch):
    from promloader import notify
    monkeypatch.setattr(notify, "recipients", lambda: [{"chat_id": "1", "events": ["order"]}])
    monkeypatch.setattr(notify, "send", lambda *a, **k: (_ for _ in ()).throw(AssertionError("не должно")))
    assert not reports.maybe_daily_summary(datetime.now().replace(hour=23))


def test_daily_summary_text_empty_day():
    assert "Заказов сегодня не было" in reports.daily_summary_text(datetime.now().date())


def test_daily_summary_event_added_to_existing_recipients():
    from promloader import notify
    with db.tx() as c:
        c.execute("INSERT INTO tg_recipients (chat_id, name, events, created_at) VALUES ('5', 'x', '[\"order\"]', ?)", (db.now(),))
    db.set_setting("tg_events_version", "1")
    assert "daily_summary" in notify.recipients()[0]["events"]
