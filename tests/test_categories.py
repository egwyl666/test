"""Категории каталога Prom: что стоит на Prom, что нужно нам, правило группы, 🔒, журнал, фид, фильтры."""

import json

import pytest

from promloader import categories, changes, db, feed, products, promcatalog

INVERT, POWERBANK = 5140401, 71109


def synced(pid):
    """Товар как будто уже выгружен: без неотправленных правок."""
    with db.tx() as c:
        c.execute("UPDATE products SET status = 'synced', synced_at = ?, pending_fields = '[]' WHERE id = ?",
                  (db.now(), pid))


def row(pid):
    return db.query_one("SELECT * FROM products WHERE id = ?", (pid,))


def test_group_rule_applies_to_unlocked_products_and_queues_synced():
    a = products.create({"name": "Инвертор 1", "price": 100, "group_name": "Инверторы"})
    b = products.create({"name": "Инвертор 2", "price": 100, "group_name": "Инверторы"})
    own = products.create({"name": "Инвертор 3", "price": 100, "group_name": "Инверторы"})
    other = products.create({"name": "Фонарь", "price": 100, "group_name": "Фонари"})
    products.update(own, {"target_category_id": POWERBANK})  # своя категория — 🔒
    for pid in (a, b, own):
        synced(pid)

    assert categories.set_group("Инверторы", INVERT) == 2
    assert row(a)["target_category_id"] == INVERT and row(b)["target_category_id"] == INVERT
    assert row(own)["target_category_id"] == POWERBANK  # закреплённую правило не трогает
    assert row(other)["target_category_id"] is None
    # выгруженный товар снова ждёт отправки, и полным импортом: категории нет среди быстрых полей
    assert row(a)["status"] == "ready" and "target_category_id" in json.loads(row(a)["pending_fields"])
    assert row(own)["status"] == "synced"
    log = db.query("SELECT * FROM product_changes WHERE product_id = ? AND field = 'target_category_id'", (a,))
    assert log[-1]["new"].endswith(f"[{INVERT}]") and log[-1]["source"] == "Категории Prom"


def test_new_product_and_group_move_follow_rule_unless_locked():
    categories.set_group("Инверторы", INVERT)
    categories.set_group("Повербанки", POWERBANK)
    pid = products.create({"name": "Инвертор", "price": 100, "group_name": "Инверторы"})
    assert row(pid)["target_category_id"] == INVERT
    products.update(pid, {"group_name": "Повербанки"})
    assert row(pid)["target_category_id"] == POWERBANK
    products.update(pid, {"group_name": "Без правила"})
    assert row(pid)["target_category_id"] is None  # нет правила — пусть решает Prom

    products.update(pid, {"target_category_id": INVERT})
    assert "target_category_id" in json.loads(row(pid)["locked_fields"])  # 🔒 даже у товара без поставщика
    products.update(pid, {"group_name": "Повербанки"})
    assert row(pid)["target_category_id"] == INVERT
    products.unlock(pid, ["target_category_id"])
    assert row(pid)["target_category_id"] == POWERBANK  # 🔒 сняли — снова по правилу группы


def test_rule_removed_and_bad_category():
    pid = products.create({"name": "Инвертор", "price": 100, "group_name": "Инверторы"})
    categories.set_group("Инверторы", INVERT)
    categories.set_group("Инверторы", None)
    assert row(pid)["target_category_id"] is None
    assert db.query_one("SELECT COUNT(*) AS n FROM group_categories")["n"] == 0
    with pytest.raises(products.ProductError):
        categories.set_group("Инверторы", "ерунда")
    assert products.parse_category("https://prom.ua/ua/c500901") == 500901


def test_preview_does_not_change_anything():
    pid = products.create({"name": "Инвертор", "price": 100, "group_name": "Инверторы"})
    synced(pid)
    before = db.query_one("SELECT COUNT(*) AS n FROM product_changes")["n"]
    p = categories.preview_group("Инверторы", INVERT)
    assert p["changed"] == 1 and p["fields"] == {"Категория Prom": 1}
    assert row(pid)["target_category_id"] is None and row(pid)["status"] == "synced"
    assert db.query_one("SELECT COUNT(*) AS n FROM group_categories")["n"] == 0
    assert db.query_one("SELECT COUNT(*) AS n FROM product_changes")["n"] == before


def test_catalog_load_records_prom_category_but_not_target():
    counts = {k: 0 for k in ("seen", "created", "updated", "skipped", "no_external_id", "kept_local", "missing_on_prom")}
    promcatalog._upsert({"id": 777, "external_id": "A-1", "name": "Инвертор", "price": 100, "presence": "available",
                         "status": "on_display", "category": {"id": INVERT, "caption": "Инверторы"}}, counts)
    pid = products.find_by_external_id("A-1")
    assert row(pid)["prom_category_id"] == INVERT and row(pid)["target_category_id"] is None
    assert categories.info(INVERT)["caption"] == "Инверторы"
    # товар с неотправленными правками тоже узнаёт, что стоит на Prom
    products.update(pid, {"target_category_id": POWERBANK})
    promcatalog._upsert({"id": 777, "external_id": "A-1", "name": "Инвертор", "status": "on_display",
                         "category": {"id": 71110, "caption": "Другое"}}, counts)
    assert row(pid)["prom_category_id"] == 71110 and row(pid)["target_category_id"] == POWERBANK
    assert categories.for_product(products.get(pid))["mismatch"]


def test_overview_split_and_mismatch():
    a = products.create({"name": "Инвертор", "price": 100, "group_name": "Инверторы"})
    b = products.create({"name": "Инвертор 2", "price": 100, "group_name": "Инверторы"})
    with db.tx() as c:
        categories.remember(c, INVERT, "Инверторы")
        categories.remember(c, POWERBANK, "Повербанки")
        c.execute("UPDATE products SET prom_category_id = ? WHERE id = ?", (INVERT, a))
        c.execute("UPDATE products SET prom_category_id = ? WHERE id = ?", (POWERBANK, b))
    categories.set_group("Инверторы", INVERT)
    g = next(x for x in categories.overview() if x["group"] == "Инверторы")
    assert g["split"] and g["mismatch"] == 1 and g["rule"]["caption"] == "Инверторы"
    assert {c["id"]: c["count"] for c in g["prom"]} == {INVERT: 1, POWERBANK: 1}
    ids = products.ids_for_filter({"cat_mismatch": "1"})
    assert ids == [b]
    assert products.ids_for_filter({"prom_category": str(INVERT)}) == [a]


def test_revert_category_change():
    pid = products.create({"name": "Инвертор", "price": 100, "group_name": "Инверторы"})
    categories.set_group("Инверторы", INVERT)
    products.update(pid, {"target_category_id": POWERBANK})
    ch = db.query("SELECT * FROM product_changes WHERE product_id = ? AND field = 'target_category_id' ORDER BY id",
                  (pid,))[-1]
    assert ch["old"].endswith(f"[{INVERT}]")
    changes.revert(ch["id"])
    assert row(pid)["target_category_id"] == INVERT


def test_feed_category_tag(monkeypatch):
    pid = products.create({"name": "Инвертор", "price": 100, "group_name": "Инверторы"})
    categories.set_group("Инверторы", INVERT)
    monkeypatch.setattr(feed, "CATEGORY_TAG", None)
    assert str(INVERT).encode() not in feed.build([pid], "")
    monkeypatch.setattr(feed, "CATEGORY_TAG", "portal_category_id")
    assert f"<portal_category_id>{INVERT}</portal_category_id>".encode() in feed.build([pid], "")


SEARCH_HTML = ('..."topHitsCategory":{"id":150403,"__typename":"Category","path":[{"id":0,"caption":"Товари",'
               '"__typename":"Category"},{"id":15,"caption":"Товари для дому","__typename":"Category"},{"id":150403,'
               '"caption":"Святкове освітлення","__typename":"Category"}]},... "category":{"id":150403},'
               '"category":{"id":150401},"category":{"id":150403} {"id":150401,"caption":"Вуличне освітлення",'
               '"__typename":"Category"}')
CATEGORY_HTML = ('<script type="application/ld+json">{"@context":"https://schema.org","@type":"BreadcrumbList",'
                 '"itemListElement":[{"@type":"ListItem","item":"https://prom.ua/ua/consumer-goods","name":"Каталог товарів"},'
                 '{"@type":"ListItem","item":"https://prom.ua/ua/Tehnika","name":"Техніка та електроніка"},'
                 '{"@type":"ListItem","item":"https://prom.ua/ua/Zaryadnye-stantsii","name":"Зарядні станції"}]}</script>'
                 '"categoryId":500901 "categoryId":500901 "categoryId":12')


def test_suggest_and_url(monkeypatch, client):
    monkeypatch.setattr(categories, "_get", lambda url, **params: SEARCH_HTML if "search" in url else CATEGORY_HTML)
    items = client.get("/api/categories/suggest", params={"q": "гірлянда"}).json()["items"]
    assert items[0]["id"] == 150403 and items[0]["path"] == "Товари для дому › Святкове освітлення"
    assert {i["id"] for i in items} == {150403, 150401}
    c = client.post("/api/categories/url", json={"url": "https://prom.ua/ua/Zaryadnye-stantsii"}).json()
    assert c["id"] == 500901 and c["caption"] == "Зарядні станції" and c["path"] == "Техніка та електроніка › Зарядні станції"
    r = client.post("/api/categories/url", json={"url": "https://example.com/x"})
    assert r.status_code == 400 and "prom.ua" in r.json()["detail"]


def test_api_group_rule_and_page(client):
    pid = products.create({"name": "Инвертор", "price": 100, "group_name": "Инверторы"})
    p = client.post("/api/categories/group/preview", json={"group": "Инверторы", "category_id": INVERT}).json()
    assert p["changed"] == 1
    r = client.post("/api/categories/group", json={"group": "Инверторы", "category_id": INVERT}).json()
    assert r["changed"] == 1 and r["groups"][0]["rule"]["id"] == INVERT
    assert row(pid)["target_category_id"] == INVERT
    body = client.get("/api/categories").json()
    assert body["groups"][0]["group"] == "Инверторы" and "send" in body
    assert client.post("/api/categories/group", json={"group": "Инверторы", "category_id": "x"}).status_code == 400
    assert client.get("/categories").status_code == 200
    assert client.get(f"/api/products/{pid}").json()["categories"]["target"]["id"] == INVERT
    assert any(c["id"] == INVERT for c in client.get("/api/meta").json()["prom_categories"])


def test_mismatch_file_for_prom_cabinet(client):
    def make(name, prom_id, prom_cat, target, status=None):
        pid = products.create({"name": name, "price": 100, "external_id": f"EXT-{prom_id or name}"})
        with db.tx() as c:
            c.execute("UPDATE products SET prom_id = ?, prom_category_id = ?, target_category_id = ? WHERE id = ?",
                      (prom_id, prom_cat, target, pid))
            if status:
                c.execute("UPDATE products SET status = ? WHERE id = ?", (status, pid))
            categories.remember(c, INVERT, "Инверторы", url="https://prom.ua/ua/Invertory")
            categories.remember(c, POWERBANK, "Повербанки")
        return pid

    make("Инвертор не там", 501, POWERBANK, INVERT)      # попадает
    make("Повербанк не там", 502, INVERT, POWERBANK)     # попадает; ссылки на категорию не знаем — пусто
    make("Инвертор на месте", 503, INVERT, INVERT)       # совпадает
    make("Нет на Prom", None, POWERBANK, INVERT)         # без номера Prom
    make("Удаляется", 505, POWERBANK, INVERT, status="deleting")
    make("Prom ещё не сказал", 506, None, INVERT)

    assert client.get("/api/categories").json()["mismatch_total"] == 2
    r = client.get("/api/categories/mismatch.xlsx")
    assert r.status_code == 200 and "kategorii-dlya-kabineta-" in r.headers["content-disposition"]
    import io

    import openpyxl
    rows = list(openpyxl.load_workbook(io.BytesIO(r.content)).active.values)
    assert list(rows[0]) == categories.MISMATCH_HEADERS
    assert rows[1] == ("501", "EXT-501", "Инвертор не там", str(INVERT), "https://prom.ua/ua/Invertory",
                       f"Инверторы [{INVERT}]", f"Повербанки [{POWERBANK}]", "https://my.prom.ua/cms/product/edit/501")
    assert rows[2][0] == "502" and rows[2][4] is None
    assert len(rows) == 3


def test_mismatch_file_empty_has_headers():
    import io

    import openpyxl
    rows = list(openpyxl.load_workbook(io.BytesIO(categories.mismatch_xlsx())).active.values)
    assert rows == [tuple(categories.MISMATCH_HEADERS)]


def test_fix_categories_trial_then_rest(monkeypatch):
    import asyncio
    import io

    import openpyxl

    ids = []
    for i in range(3):
        pid = products.create({"name": f"Инвертор {i}", "price": 100, "external_id": f"INV-{i}"})
        with db.tx() as c:
            c.execute("UPDATE products SET prom_id = ?, prom_category_id = ?, target_category_id = ? WHERE id = ?",
                      (700 + i, POWERBANK, INVERT, pid))
        ids.append(pid)
    live = {700 + i: {"id": 700 + i, "name": f"Инвертор {i}", "sku": f"S{i}", "external_id": f"INV-{i}", "price": 100,
                      "presence": "available", "status": "on_display", "category": {"id": POWERBANK}} for i in range(3)}
    files = []

    class Fake:
        async def list_products(self, limit=100, last_id=None):
            return {"products": [] if last_id else [dict(p, category=dict(p["category"])) for p in live.values()]}

        async def import_file(self, content, settings, name, ctype):
            rows = list(openpyxl.load_workbook(io.BytesIO(content)).active.values)
            files.append(rows)
            assert settings == promcatalog.EXT_IMPORT_SETTINGS
            for r in rows[1:]:
                live[int(r[0])]["category"] = {"id": int(r[4])}  # Prom меняет только категорию
            return "imp"

        async def import_status(self, iid):
            return {"status": "SUCCESS", "total": 1, "imported": 1, "updated": 1}

    result = asyncio.run(promcatalog.fix_categories(Fake()))
    assert result == {"done": 3, "skipped": 0}
    assert [len(f) - 1 for f in files] == [2, 1]  # проба на 2, потом остальные
    assert files[0][0] == tuple(promcatalog.CAT_HEADERS) and files[0][1] == ("700", "Инвертор 0", "S0", "INV-0", str(INVERT))
    assert all(row(pid)["prom_category_id"] == INVERT for pid in ids)
    assert categories.mismatch_rows() == []
