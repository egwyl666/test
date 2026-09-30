"""Массовый импорт из Excel/CSV: пользователь выбирает лист, строки и что лежит в каждой колонке."""

import csv
import io
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from . import products

# Куда можно направить колонку. "param" — характеристика с именем из заголовка колонки.
TARGETS = {
    "": "— не импортировать —",
    "external_id": "Артикул / код",
    "name": "Название",
    "name_ua": "Название (укр.)",
    "price": "Цена",
    "old_price": "Старая цена (до скидки)",
    "currency": "Валюта",
    "quantity": "Количество",
    "presence": "Наличие",
    "group_name": "Группа / категория",
    "vendor": "Производитель",
    "country": "Страна",
    "description": "Описание",
    "description_ua": "Описание (укр.)",
    "keywords": "Ключевые слова",
    "unit": "Единица измерения",
    "images": "Фото (ссылки)",
    "param": "Характеристика",
}

# Подсказки для автоматического сопоставления по заголовкам колонок.
HINTS = [
    ("name_ua", r"назв\w*.*(укр|ua)|найменування"),
    ("description_ua", r"опис\w*.*(укр|ua)|^опис$"),
    ("old_price", r"стар\w* цен|цена до|стара ціна|old.?price"),
    ("external_id", r"артикул|код|sku|external|ідентиф|идентиф"),
    ("name", r"назв|наимен|товар|name"),
    ("price", r"цен|ціна|price|стоим|вартість"),
    ("currency", r"валют|currency"),
    ("quantity", r"кол-?во|колич|кільк|остат|залиш|qty|quantity|stock"),
    ("presence", r"налич|наявн|presence|available"),
    ("group_name", r"групп|груп|катег|розділ|раздел|category|group"),
    ("vendor", r"произв|виробн|бренд|brand|vendor|марка"),
    ("country", r"стран|країн|country"),
    ("description", r"опис|описан|description"),
    ("keywords", r"ключ|keyword|теги|tags"),
    ("unit", r"един|одиниц|unit"),
    ("images", r"фото|изображ|зображ|картин|image|photo|picture"),
]

MAX_PREVIEW_ROWS = 2000


class ImportError_(ValueError):
    pass


# ---------- чтение файла ----------

def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M") if value.time() != datetime.min.time() else value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _read_csv(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportError_("Не удалось определить кодировку CSV — сохраните файл в UTF-8")
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def _load(path: Path):
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        raise ImportError_(f"Не удалось открыть файл как Excel (.xlsx): {exc}")


def sheet_names(path: Path) -> list[str]:
    if path.suffix.lower() == ".csv":
        return ["CSV"]
    return _load(path).sheetnames


def read_sheet(path: Path, sheet: str) -> tuple[list[list[str]], dict[int, list[bytes]]]:
    """Все строки листа как текст + картинки, вставленные в ячейки: {номер строки: [байты]}."""
    if path.suffix.lower() == ".csv":
        return _read_csv(path), {}
    wb = _load(path)
    if sheet not in wb.sheetnames:
        raise ImportError_(f"В файле нет листа «{sheet}»")
    ws = wb[sheet]
    rows = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    # хвост из пустых строк Excel любит оставлять — отрезаем
    while rows and not any(rows[-1]):
        rows.pop()
    images: dict[int, list[bytes]] = {}
    for img in getattr(ws, "_images", []):
        try:
            row_number = img.anchor._from.row + 1
            images.setdefault(row_number, []).append(img._data())
        except Exception:
            continue
    return rows, images


# ---------- выбор строк ----------

def parse_row_spec(spec: str, max_row: int) -> list[int]:
    """'2-50, 55, 60-' -> номера строк (как в Excel, с единицы)."""
    spec = (spec or "").strip()
    if not spec:
        return []
    selected: set[int] = set()
    for part in re.split(r"[,;\s]+", spec):
        if not part:
            continue
        m = re.fullmatch(r"(\d+)?\s*-\s*(\d+)?|(\d+)", part)
        if not m:
            raise ImportError_(f"Не понял диапазон «{part}». Пример: 2-50, 55, 60-")
        if m.group(3):
            start = end = int(m.group(3))
        else:
            start = int(m.group(1) or 1)
            end = int(m.group(2) or max_row)
        if start > end:
            start, end = end, start
        selected.update(range(max(start, 1), min(end, max_row) + 1))
    return sorted(selected)


# ---------- сопоставление колонок ----------

def guess_mapping(headers: list[str]) -> dict[str, str]:
    mapping, used = {}, set()
    for idx, header in enumerate(headers):
        text = (header or "").strip().lower()
        if not text:
            continue
        for target, pattern in HINTS:
            if target not in used and re.search(pattern, text):
                mapping[get_column_letter(idx + 1)] = target
                used.add(target)
                break
    return mapping


def _split_urls(text: str) -> list[str]:
    return [u for u in re.split(r"[\s,;|]+", text or "") if re.match(r"^https?://", u)]


def build_products(
    rows: list[list[str]],
    header_row: int,
    row_numbers: list[int],
    mapping: dict[str, str],
    defaults: dict | None = None,
    embedded_images: dict[int, list[bytes]] | None = None,
) -> list[dict]:
    """Строки -> товары. Для каждой строки: данные, фото, ошибки и предупреждения."""
    defaults = {k: v for k, v in (defaults or {}).items() if v not in (None, "")}
    headers = rows[header_row - 1] if 0 < header_row <= len(rows) else []
    columns = []
    for letter, target in mapping.items():
        if not target:
            continue
        if target not in TARGETS and not target.startswith("param:"):
            raise ImportError_(f"Неизвестное поле «{target}»")
        idx = openpyxl.utils.column_index_from_string(letter) - 1
        if target == "param":
            header = headers[idx] if idx < len(headers) else ""
            target = "param:" + (header or f"Колонка {letter}")
        columns.append((idx, target))

    result = []
    for number in row_numbers:
        if number < 1 or number > len(rows) or number == header_row:
            continue
        row = rows[number - 1]
        if not any(row):
            continue
        raw, params, urls, errors = dict(defaults), [], [], []
        for idx, target in columns:
            value = row[idx] if idx < len(row) else ""
            if target.startswith("param:"):
                if value:
                    params.append({"name": target[6:], "value": value})
            elif target == "images":
                urls += _split_urls(value)
            elif value != "":
                raw[target] = value
        if params:
            raw["params"] = params
        try:
            data = products.normalize(raw)
        except products.ProductError as exc:
            errors.append(str(exc))
            data = {}
        # наличие по количеству, если колонки «наличие» нет
        if "presence" not in raw and data.get("quantity") is not None:
            data["presence"] = "available" if data["quantity"] > 0 else "not_available"
        files = (embedded_images or {}).get(number, [])
        check = products.validate({**data, "params": params}, image_count=len(urls) + len(files))
        result.append({
            "row": number,
            "data": data,
            "params": params,
            "image_urls": urls[: products.MAX_IMAGES],
            "embedded_images": len(files),
            "errors": errors + check["errors"],
            "warnings": check["warnings"],
        })
    return result


def headers_for(rows: list[list[str]], header_row: int) -> list[str]:
    width = max((len(r) for r in rows), default=0)
    headers = rows[header_row - 1] if 0 < header_row <= len(rows) else []
    return [headers[i] if i < len(headers) else "" for i in range(width)]


def column_letters(width: int) -> list[str]:
    return [get_column_letter(i + 1) for i in range(width)]
