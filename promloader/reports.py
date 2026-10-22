"""Отчёты: продажи, товары, поставщики и группы, состояние магазина — из того, что уже есть в программе.

Источник продаж — заказы Prom (таблица `orders`, ответ Prom целиком в `data`). Отменённые заказы в выручку,
штуки, средний чек и прибыль не входят — для них отдельно доля отмен и причины.

Прибыль позиции = сумма позиции − закупка в гривнах на дату заказа × количество. Комиссия Prom (`cpa_commission`)
вычитается из прибыли заказа отдельной строкой — она есть почти у каждого заказа и заметно меняет картину.
Закупку на дату заказа восстанавливаем по журналу изменений: первое изменение закупки после заказа хранит, какой
она была до него. Нет ни одной записи о закупке — берём текущую и помечаем «≈». Валюта — по курсу программы
сейчас (истории курса программа не хранит).

Товары заказа связываются с программой по номеру товара на Prom (`products.prom_id`), затем по артикулу:
внешний ID в заказах Prom часто пустой.
"""

import io
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone

from . import db, products, rates

PERIODS = {"today": "Сегодня", "7d": "7 дней", "30d": "30 дней", "month": "Этот месяц", "custom": "Свой период"}
DEAD_DAYS = 30            # «лежат без продаж» по умолчанию
COST_FIELDS = ("cost_price", "cost_currency")


class ReportError(ValueError):
    pass


# ---------- период ----------

def _local(iso: str) -> datetime | None:
    """Время заказа Prom (ISO, обычно с поясом) → местное время компьютера."""
    if not iso:
        return None
    try:
        dt = datetime.fromisoformat(str(iso).replace("Z", "+00:00"))
    except ValueError:
        return None
    return dt.astimezone() if dt.tzinfo else dt


def period_bounds(period: str = "30d", date_from: str = "", date_to: str = "", today: date | None = None) -> tuple[date, date]:
    """Первый и последний день периода включительно (местные даты)."""
    today = today or datetime.now().date()
    if period == "today":
        return today, today
    if period == "7d":
        return today - timedelta(days=6), today
    if period == "month":
        return today.replace(day=1), today
    if period == "custom":
        try:
            start = date.fromisoformat(date_from) if date_from else today - timedelta(days=29)
            end = date.fromisoformat(date_to) if date_to else today
        except ValueError:
            raise ReportError("Дата — в виде ГГГГ-ММ-ДД") from None
        if start > end:
            start, end = end, start
        return start, end
    return today - timedelta(days=29), today


# ---------- заказы ----------

def _num(value) -> float:
    try:
        return products.parse_number(value) or 0.0
    except products.ProductError:
        return 0.0


def _load_orders(start: date | None = None, end: date | None = None) -> list[dict]:
    """Заказы периода (по местной дате создания) в удобном виде."""
    out = []
    # date_created в базе — строка Prom; грубо отсекаем по строке с запасом в сутки (пояс), точно — ниже
    where, args = "", []
    if start:
        where = "WHERE date_created >= ?"
        args.append((start - timedelta(days=1)).isoformat())
    for row in db.query(f"SELECT id, status, date_created, data FROM orders {where} ORDER BY date_created", args):
        when = _local(row["date_created"])
        if when is None or (start and when.date() < start) or (end and when.date() > end):
            continue
        data = json.loads(row["data"] or "{}")
        commission = data.get("cpa_commission") or {}
        cancellation = data.get("cancellation") or {}
        items = []
        for p in data.get("products") or []:
            if not isinstance(p, dict):
                continue
            qty = _num(p.get("quantity")) or 1.0
            total = _num(p.get("total_price")) or _num(p.get("price")) * qty
            items.append({"prom_id": p.get("id"), "external_id": p.get("external_id") or "", "sku": p.get("sku") or "",
                          "name": p.get("name") or (p.get("name_multilang") if isinstance(p.get("name_multilang"), dict) else {}).get("ru") or "",
                          "qty": qty, "total": total})
        out.append({
            "id": row["id"], "status": row["status"], "when": when, "day": when.date(),
            "total": _num(data.get("price")) or sum(i["total"] for i in items),
            "commission": 0.0 if (commission.get("is_refunded") if isinstance(commission, dict) else False)
            else _num(commission.get("amount") if isinstance(commission, dict) else commission),
            "canceled": row["status"] == "canceled",
            "cancel_reason": (cancellation.get("title") if isinstance(cancellation, dict) else "") or "",
            "items": items,
        })
    return out


def _link_products(orders: list[dict]) -> dict:
    """Позиция заказа → товар программы. Одним запросом на все позиции, без запроса на каждую."""
    prom_ids = {int(i["prom_id"]) for o in orders for i in o["items"] if str(i["prom_id"] or "").isdigit()}
    codes = {i["external_id"] for o in orders for i in o["items"] if i["external_id"]} | \
            {i["sku"] for o in orders for i in o["items"] if i["sku"]}
    rows = db.query(f"""SELECT p.id, p.name, p.prom_id, p.external_id, p.group_name, p.supplier_id, p.cost_price,
                               p.cost_currency, p.currency, p.created_at, s.name AS supplier_name
                        FROM products p LEFT JOIN suppliers s ON s.id = p.supplier_id
                        WHERE p.prom_id {db.IN_LIST} OR p.external_id {db.IN_LIST}""",
                    (db.as_list(sorted(prom_ids)), db.as_list(sorted(codes))))
    by_prom = {r["prom_id"]: dict(r) for r in rows if r["prom_id"]}
    by_code = {r["external_id"]: dict(r) for r in rows}
    for o in orders:
        for i in o["items"]:
            i["product"] = (by_prom.get(int(i["prom_id"])) if str(i["prom_id"] or "").isdigit() else None) \
                or by_code.get(i["external_id"]) or by_code.get(i["sku"])
    return {p["id"]: p for p in list(by_prom.values()) + list(by_code.values())}


def _cost_history(product_ids) -> dict[int, list]:
    """Изменения закупки по товарам: {id: [(время, поле, было), …]} по времени — одним запросом."""
    hist: dict[int, list] = defaultdict(list)
    if not product_ids:
        return hist
    for r in db.query(f"""SELECT product_id, at, field, old FROM product_changes
                          WHERE product_id {db.IN_LIST} AND field IN ('cost_price', 'cost_currency') ORDER BY at, id""",
                      (db.as_list(sorted(product_ids)),)):
        hist[r["product_id"]].append((r["at"], r["field"], r["old"]))
    return hist


def _cost_at(p: dict, hist: list, when: datetime) -> tuple[float | None, str, bool]:
    """Закупка и её валюта на момент заказа; третье — «≈» (точной истории нет, взята текущая)."""
    value = {"cost_price": p["cost_price"], "cost_currency": p["cost_currency"] or p["currency"] or "UAH"}
    # журнал пишет время в UTC (db.now) — сравниваем в том же виде
    at = (when if when.tzinfo else when.astimezone()).astimezone(timezone.utc).isoformat(timespec="seconds")
    for field in COST_FIELDS:
        later = [old for t, f, old in hist if f == field and t > at]
        if later:  # первое изменение после заказа знает, какое значение было в момент заказа
            old = later[0]
            if field == "cost_price":
                try:
                    value[field] = float(old) if old not in ("", None) else None
                except ValueError:
                    value[field] = None
            else:
                value[field] = old or value[field]
    approx = not any(f == "cost_price" for _, f, _ in hist)
    return value["cost_price"], (value["cost_currency"] or "UAH").upper(), approx


class _Money:
    """Закупка → гривны по курсу программы (один Table на отчёт)."""

    def __init__(self):
        self.table = rates.Table()

    def uah(self, amount, currency, supplier_id=None):
        if amount is None:
            return None
        try:
            rate = self.table.rate(currency, supplier_id)
        except rates.RateError:
            return None
        return amount * rate if rate else amount


def _profit_items(orders: list[dict]) -> None:
    """У каждой позиции: закупка (грн) на дату заказа, прибыль, «≈»."""
    _link_products(orders)
    ids = {i["product"]["id"] for o in orders for i in o["items"] if i["product"]}
    hist = _cost_history(ids)
    money = _Money()
    for o in orders:
        for i in o["items"]:
            i["cost"] = i["profit"] = None
            i["approx"] = False
            p = i["product"]
            if not p:
                continue
            cost, currency, approx = _cost_at(p, hist.get(p["id"], []), o["when"])
            cost_uah = money.uah(cost, currency, p["supplier_id"])
            if cost_uah is None:
                continue
            i["cost"] = round(cost_uah * i["qty"], 2)
            i["profit"] = round(i["total"] - i["cost"], 2)
            i["approx"] = approx


def _r(x):
    return round(x, 2) if isinstance(x, (int, float)) else x


# ---------- 1. продажи ----------

def sales(start: date, end: date) -> dict:
    orders = _load_orders(start, end)
    _profit_items(orders)
    done = [o for o in orders if not o["canceled"]]
    revenue = sum(o["total"] for o in done)
    known = [i for o in done for i in o["items"] if i["profit"] is not None]
    items_total = sum(i["total"] for o in done for i in o["items"])
    commission = sum(o["commission"] for o in done)
    profit = sum(i["profit"] for i in known) - sum(o["commission"] for o in done if any(i["profit"] is not None for i in o["items"]))
    days = []
    by_day = defaultdict(lambda: {"orders": 0, "revenue": 0.0, "profit": 0.0, "canceled": 0})
    for o in orders:
        d = by_day[o["day"]]
        if o["canceled"]:
            d["canceled"] += 1
            continue
        d["orders"] += 1
        d["revenue"] += o["total"]
        d["profit"] += sum(i["profit"] for i in o["items"] if i["profit"] is not None) - \
            (o["commission"] if any(i["profit"] is not None for i in o["items"]) else 0)
    day = start
    while day <= end:
        v = by_day[day]
        days.append({"date": day.isoformat(), "orders": v["orders"], "revenue": _r(v["revenue"]),
                     "profit": _r(v["profit"]), "canceled": v["canceled"]})
        day += timedelta(days=1)
    canceled = [o for o in orders if o["canceled"]]
    reasons = Counter(o["cancel_reason"] or "Причина не указана" for o in canceled)
    return {
        "orders": len(done), "revenue": _r(revenue), "avg_check": _r(revenue / len(done)) if done else 0,
        "units": sum(i["qty"] for o in done for i in o["items"]),
        "profit": _r(profit) if known else None,
        # какую часть выручки покрывает прибыль: у остальных позиций закупка неизвестна
        "profit_coverage": round(sum(i["total"] for i in known) / items_total * 100) if items_total else 0,
        "approx": any(i["approx"] for i in known), "commission": _r(commission),
        "canceled": len(canceled), "cancel_rate": round(len(canceled) / len(orders) * 100, 1) if orders else 0,
        "cancel_reasons": [{"reason": k, "count": n} for k, n in reasons.most_common()],
        "statuses": dict(Counter(o["status"] for o in orders)),
        "days": days,
    }


# ---------- 2. товары ----------

def _abc(rows: list[dict]) -> None:
    """A — товары, которые дают первые 80% выручки, B — следующие 15%, C — остальные 5%."""
    total = sum(r["revenue"] for r in rows) or 1
    acc = 0.0
    for r in sorted(rows, key=lambda r: -r["revenue"]):
        share_before = acc / total
        acc += r["revenue"]
        r["abc"] = "A" if share_before < 0.80 else "B" if share_before < 0.95 else "C"


def products_report(start: date, end: date, dead_days: int = DEAD_DAYS) -> dict:
    orders = _load_orders(start, end)
    _profit_items(orders)
    rows = _product_rows(orders)
    sold = [r for r in rows if r["units"]]
    top = sorted(sold, key=lambda r: (-r["revenue"], r["name"]))
    return {
        "top": top,
        "abc": {k: sum(1 for r in sold if r.get("abc") == k) for k in "ABC"},
        "dead": dead_stock(dead_days),
        "dead_days": dead_days,
        "canceled": sorted([r for r in rows if r["canceled"]], key=lambda r: -r["canceled"])[:30],
    }


def _product_rows(orders: list[dict]) -> list[dict]:
    """Позиции заказов (уже с прибылью) → строка на товар: штуки, выручка, прибыль, заказы, отмены, ABC."""
    rows: dict = {}
    canceled: Counter = Counter()
    for o in orders:
        for i in o["items"]:
            p = i["product"]
            key = ("p", p["id"]) if p else ("prom", i["prom_id"] or i["name"])
            r = rows.setdefault(key, {"product_id": p["id"] if p else None, "name": (p or {}).get("name") or i["name"],
                                      "code": (p or {}).get("external_id") or i["sku"] or "",
                                      "group": (p or {}).get("group_name") or "", "supplier": (p or {}).get("supplier_name") or "",
                                      "units": 0.0, "revenue": 0.0, "profit": 0.0, "profit_known": False,
                                      "approx": False, "orders": 0, "canceled": 0})
            if o["canceled"]:
                r["canceled"] += 1
                canceled[key] += 1
                continue
            r["units"] += i["qty"]
            r["revenue"] += i["total"]
            r["orders"] += 1
            if i["profit"] is not None:
                r["profit"] += i["profit"]
                r["profit_known"] = True
                r["approx"] = r["approx"] or i["approx"]
    sold = [r for r in rows.values() if r["units"]]
    _abc(sold)
    for r in rows.values():
        r["revenue"], r["profit"] = _r(r["revenue"]), _r(r["profit"]) if r["profit_known"] else None
        r["units"] = _r(r["units"])
    return list(rows.values())


def dead_stock(days: int = DEAD_DAYS) -> list[dict]:
    """В наличии и на Prom, но без продаж N дней (или ни разу, если товар старше N дней)."""
    since = datetime.now().date() - timedelta(days=days)
    orders = _load_orders()
    _link_products(orders)
    last_sale: dict[int, date] = {}
    for o in orders:
        if o["canceled"]:
            continue
        for i in o["items"]:
            if i["product"]:
                pid = i["product"]["id"]
                last_sale[pid] = max(last_sale.get(pid, o["day"]), o["day"])
    out = []
    for p in db.query("""SELECT p.id, p.name, p.external_id, p.group_name, p.price, p.created_at, s.name AS supplier_name
                         FROM products p LEFT JOIN suppliers s ON s.id = p.supplier_id
                         WHERE p.presence = 'available' AND p.synced_at IS NOT NULL ORDER BY p.name"""):
        sold = last_sale.get(p["id"])
        created = (p["created_at"] or "")[:10]
        if (sold and sold < since) or (not sold and created and created < since.isoformat()):
            out.append({"product_id": p["id"], "name": p["name"], "code": p["external_id"], "group": p["group_name"],
                        "supplier": p["supplier_name"] or "", "price": p["price"],
                        "last_sale": sold.isoformat() if sold else ""})
    return out


# ---------- 3. поставщики и группы ----------

def suppliers_report(start: date, end: date) -> dict:
    orders = _load_orders(start, end)
    _profit_items(orders)

    def agg(key_fn):
        acc: dict = {}
        for o in orders:
            if o["canceled"]:
                continue
            for i in o["items"]:
                k = key_fn(i)
                r = acc.setdefault(k, {"name": k, "units": 0.0, "revenue": 0.0, "profit": 0.0, "profit_known": False,
                                       "approx": False})
                r["units"] += i["qty"]
                r["revenue"] += i["total"]
                if i["profit"] is not None:
                    r["profit"] += i["profit"]
                    r["profit_known"] = True
                    r["approx"] = r["approx"] or i["approx"]
        rows = sorted(acc.values(), key=lambda r: -r["revenue"])
        for r in rows:
            r["revenue"], r["profit"] = _r(r["revenue"]), _r(r["profit"]) if r["profit_known"] else None
        return rows

    stock = {r["id"]: dict(r) for r in db.query("""
        SELECT s.id, s.name,
               (SELECT COUNT(*) FROM products p WHERE p.supplier_id = s.id) AS products,
               (SELECT COUNT(*) FROM products p WHERE p.supplier_id = s.id AND p.synced_at IS NOT NULL) AS on_prom,
               (SELECT COUNT(*) FROM supplier_items si WHERE si.supplier_id = s.id AND si.missing = 1
                    AND si.product_id IS NOT NULL) AS missing
        FROM suppliers s ORDER BY s.name""")}
    sales_by_supplier = {r["name"]: r for r in agg(lambda i: (i["product"] or {}).get("supplier_name") or "")}
    suppliers = []
    for s in stock.values():
        r = sales_by_supplier.pop(s["name"], {"units": 0, "revenue": 0, "profit": None, "approx": False})
        suppliers.append({**s, **{k: r[k] for k in ("units", "revenue", "profit", "approx")}})
    rest = sales_by_supplier.pop("", None)
    if rest:  # проданное без поставщика (свои товары или не найденные в программе)
        suppliers.append({"id": None, "name": "", "products": None, "on_prom": None, "missing": None,
                          **{k: rest[k] for k in ("units", "revenue", "profit", "approx")}})
    return {"suppliers": suppliers, "groups": agg(lambda i: (i["product"] or {}).get("group_name") or "")}


# ---------- 4. магазин ----------

def store_report(start: date, end: date) -> dict:
    from . import aiprice

    statuses = {r["status"]: r["n"] for r in db.query("SELECT status, COUNT(*) AS n FROM products GROUP BY status")}
    # только нужные колонки: полные строки списка (фото, поставщик, описание целиком) на 20 000 товаров — секунда
    errors_sql = products.ERRORS_SQL.format(min_price=products.MIN_PRICE)
    rows = db.query(f"""SELECT p.name, p.name_ua, TRIM(COALESCE(p.description, ''), {products.BLANK_SQL}) != '' AS has_description,
                              EXISTS (SELECT 1 FROM images i WHERE i.product_id = p.id) AS has_photo,
                              {errors_sql} AS has_errors FROM products p""")
    quality = {"no_photo": 0, "no_description": 0, "no_ua": 0, "errors": 0}
    for r in rows:
        quality["no_photo"] += 0 if r["has_photo"] else 1
        quality["no_description"] += 0 if r["has_description"] else 1
        if not (r["name_ua"] or "").strip() and not products.looks_ukrainian(r["name"] or ""):
            quality["no_ua"] += 1
        quality["errors"] += 1 if r["has_errors"] else 0
    first, last = start.isoformat(), (end + timedelta(days=1)).isoformat()
    ups = downs = 0
    pct = []
    for r in db.query("SELECT old, new FROM product_changes WHERE field = 'price' AND at >= ? AND at < ?", (first, last)):
        try:
            a, b = float(r["old"]), float(r["new"])
        except (TypeError, ValueError):
            continue
        if a > 0 and b != a:
            ups += b > a
            downs += b < a
            pct.append((b - a) / a * 100)
    ai_rows = db.query("""SELECT model, action, COUNT(*) AS n, SUM(tokens_in) AS tin, SUM(tokens_out) AS tout,
                                 SUM(COALESCE(cost_usd, 0)) AS usd FROM ai_usage WHERE at >= ? AND at < ?
                          GROUP BY model, action ORDER BY usd DESC""", (first, last))
    ai = [{"model": r["model"], "action": r["action"], "requests": r["n"], "tokens": (r["tin"] or 0) + (r["tout"] or 0),
           "usd": round(r["usd"] or 0, 4), "uah": aiprice.to_uah(r["usd"] or 0)} for r in ai_rows]
    return {
        "statuses": statuses, "total": len(rows), "quality": quality,
        "prices": {"changed": ups + downs, "up": ups, "down": downs,
                   "avg_pct": round(sum(pct) / len(pct), 1) if pct else 0},
        "ai": ai, "ai_usd": round(sum(a["usd"] for a in ai), 4),
        "ai_uah": aiprice.to_uah(sum(a["usd"] for a in ai)),
    }


TABS = {"sales": sales, "products": products_report, "suppliers": suppliers_report, "store": store_report}


def build(tab: str, period: str = "30d", date_from: str = "", date_to: str = "", dead_days: int = DEAD_DAYS) -> dict:
    if tab not in TABS:
        raise ReportError("Неизвестный отчёт")
    start, end = period_bounds(period, date_from, date_to)
    rates.prefetch()  # курс — до расчёта, а не посреди него
    data = products_report(start, end, dead_days) if tab == "products" else TABS[tab](start, end)
    return {"tab": tab, "period": period, "from": start.isoformat(), "to": end.isoformat(), **data}


# ---------- страница «Заказы»: выручка и статистика товаров ----------

def _totals(orders: list[dict]) -> dict:
    done = [o for o in orders if not o["canceled"]]
    revenue = sum(o["total"] for o in done)
    known = [o for o in done if any(i["profit"] is not None for i in o["items"])]
    profit = sum(i["profit"] for o in known for i in o["items"] if i["profit"] is not None) - \
        sum(o["commission"] for o in known)
    return {"orders": len(done), "revenue": _r(revenue), "avg_check": _r(revenue / len(done)) if done else 0,
            "units": _r(sum(i["qty"] for o in done for i in o["items"])),
            "profit": _r(profit) if known else None,
            "approx": any(i["approx"] for o in known for i in o["items"] if i["profit"] is not None),
            "canceled": len(orders) - len(done)}


def order_summary(today: date | None = None) -> dict:
    """Полоска над списком заказов: сегодня, вчера, 7 дней, этот месяц — один проход по заказам месяца."""
    today = today or datetime.now().date()
    start = min(today.replace(day=1), today - timedelta(days=6))
    rates.prefetch()
    orders = _load_orders(start, today)
    _profit_items(orders)
    spans = {"today": (today, today), "yesterday": (today - timedelta(days=1),) * 2,
             "7d": (today - timedelta(days=6), today), "month": (today.replace(day=1), today)}
    return {k: _totals([o for o in orders if a <= o["day"] <= b]) for k, (a, b) in spans.items()}


ORDER_SORTS = {"revenue": lambda r: (-r["revenue"], r["name"]), "units": lambda r: (-r["units"], r["name"]),
               "profit": lambda r: (-(r["profit"] if r["profit"] is not None else float("-inf")), r["name"]),
               "orders": lambda r: (-r["orders"], r["name"]), "canceled": lambda r: (-r["canceled"], r["name"]),
               "name": lambda r: r["name"].lower()}


def order_stats(period: str = "30d", date_from: str = "", date_to: str = "", q: str = "", sort: str = "revenue") -> dict:
    """Вкладка «Статистика» на странице заказов: итоги периода и строка на каждый проданный (или отменённый) товар."""
    start, end = period_bounds(period, date_from, date_to)
    rates.prefetch()
    orders = _load_orders(start, end)
    _profit_items(orders)
    rows = _product_rows(orders)
    total = sum(r["revenue"] for r in rows) or 0
    for r in rows:
        r["share"] = round(r["revenue"] / total * 100, 1) if total else 0
    if q:
        q = q.lower()
        rows = [r for r in rows if q in (r["name"] or "").lower() or q in (r["code"] or "").lower()]
    rows.sort(key=ORDER_SORTS.get(sort, ORDER_SORTS["revenue"]))
    return {"tab": "order_products", "period": period, "from": start.isoformat(), "to": end.isoformat(),
            "totals": _totals(orders), "items": rows}


# ---------- сводка за день в Telegram ----------

SUMMARY_TIME = "21:00"   # по умолчанию, местное время: к вечеру день уже почти закончен


def summary_time() -> str:
    """Время сводки из «Настроек» (ЧЧ:ММ)."""
    value = db.get_setting("daily_summary_time") or SUMMARY_TIME
    return value if re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", value) else SUMMARY_TIME


ALL_DAYS = "1,2,3,4,5,6,7"   # ISO: 1 — понедельник … 7 — воскресенье


def summary_days() -> list[int]:
    """Дни недели, в которые приходит сводка (настройка daily_summary_days)."""
    raw = db.get_setting("daily_summary_days")
    if raw == "":
        return [1, 2, 3, 4, 5, 6, 7]
    return sorted({int(x) for x in raw.split(",") if x.strip().isdigit() and 1 <= int(x) <= 7})


def clean_summary_days(value) -> str:
    """Список или строка «1,3,5» → «1,3,5». Пустой выбор допустим — сводка не приходит совсем."""
    items = value if isinstance(value, (list, tuple)) else str(value or "").split(",")
    days = set()
    for x in items:
        x = str(x).strip()
        if not x:
            continue
        if not x.isdigit() or not 1 <= int(x) <= 7:
            raise ReportError("Дни сводки — числа от 1 (понедельник) до 7 (воскресенье)")
        days.add(int(x))
    return ",".join(str(d) for d in sorted(days)) or "none"


def clean_summary_time(value: str) -> str:
    value = str(value or "").strip()
    m = re.fullmatch(r"(\d{1,2})(?::(\d{2}))?", value)
    if not m or int(m.group(1)) > 23 or int(m.group(2) or 0) > 59:
        raise ReportError("Время сводки — в виде ЧЧ:ММ, например 21:00")
    return f"{int(m.group(1)):02d}:{int(m.group(2) or 0):02d}"


def _money(x) -> str:
    return f"{x:,.0f}".replace(",", " ") + " грн"


def daily_summary_text(day: date) -> str:
    from .notify import esc
    orders = _load_orders(day - timedelta(days=1), day)
    _profit_items(orders)
    t = _totals([o for o in orders if o["day"] == day])
    y = _totals([o for o in orders if o["day"] == day - timedelta(days=1)])
    lines = [f"📊 <b>Итоги дня {day.strftime('%d.%m.%Y')}</b>"]
    if not t["orders"] and not t["canceled"]:
        lines.append("Заказов сегодня не было.")
    else:
        diff = ""
        if y["revenue"]:
            pct = round((t["revenue"] - y["revenue"]) / y["revenue"] * 100)
            diff = f" ({'+' if pct >= 0 else ''}{pct}% к вчера)"
        lines.append(f"Заказов: <b>{t['orders']}</b> · выручка <b>{_money(t['revenue'])}</b>{diff}")
        if t["orders"]:
            lines.append(f"Средний чек: {_money(t['avg_check'])}")
        if t["profit"] is not None:
            lines.append(f"Прибыль: {'≈ ' if t['approx'] else ''}{_money(t['profit'])}")
        if t["canceled"]:
            lines.append(f"Отменено: {t['canceled']}")
        top = sorted((r for r in _product_rows([o for o in orders if o["day"] == day]) if r["units"]),
                     key=lambda r: -r["revenue"])[:3]
        if top:
            lines.append("Больше всего выручки:")
            lines += [f"• {esc(r['name'])} — {r['units']:g} шт., {_money(r['revenue'])}" for r in top]
    return "\n".join(lines)


def maybe_daily_summary(now: datetime | None = None) -> bool:
    """Раз в день после времени из «Настроек» (summary_time) — сводка подписанным на «daily_summary». Пропущенный вечер (компьютер был
    выключен) не досылается утром: вчерашние итоги уже не новость."""
    from . import notify
    now = now or datetime.now()
    today = now.date().isoformat()
    if now.strftime("%H:%M") < summary_time() or db.get_setting("daily_summary_sent") == today:
        return False
    if now.isoweekday() not in summary_days():
        return False  # в этот день недели сводка выключена (не помечаем «отправлено» — день просто пропускается)
    db.set_setting("daily_summary_sent", today)
    if not any("daily_summary" in r["events"] for r in notify.recipients()):
        return False
    rates.prefetch()
    notify.send("daily_summary", daily_summary_text(now.date()), background=False)
    return True


# ---------- Excel ----------

# заголовки листов Excel (как и выгрузка списка товаров) — по-русски
XLSX_TABLES = {
    "sales": [("По дням", "days", [("Дата", "date"), ("Заказы", "orders"), ("Выручка, грн", "revenue"),
                                   ("Прибыль, грн", "profit"), ("Отменено", "canceled")]),
              ("Причины отмен", "cancel_reasons", [("Причина", "reason"), ("Заказов", "count")])],
    "products": [("Продажи товаров", "top", [("Товар", "name"), ("Артикул", "code"), ("Группа", "group"),
                                             ("Поставщик", "supplier"), ("Штук", "units"), ("Выручка, грн", "revenue"),
                                             ("Прибыль, грн", "profit"), ("≈", "approx"), ("ABC", "abc")]),
                 ("Без продаж", "dead", [("Товар", "name"), ("Артикул", "code"), ("Группа", "group"),
                                         ("Поставщик", "supplier"), ("Цена", "price"), ("Последняя продажа", "last_sale")]),
                 ("Отмены", "canceled", [("Товар", "name"), ("Артикул", "code"), ("Отмен", "canceled")])],
    "suppliers": [("Поставщики", "suppliers", [("Поставщик", "name"), ("Товаров", "products"), ("На Prom", "on_prom"),
                                               ("Пропали из прайса", "missing"), ("Штук продано", "units"),
                                               ("Выручка, грн", "revenue"), ("Прибыль, грн", "profit")]),
                  ("Группы", "groups", [("Группа", "name"), ("Штук", "units"), ("Выручка, грн", "revenue"),
                                        ("Прибыль, грн", "profit")])],
    "order_products": [("Товары в заказах", "items", [("Товар", "name"), ("Артикул", "code"), ("Группа", "group"),
                                                       ("Поставщик", "supplier"), ("Заказов", "orders"), ("Штук", "units"),
                                                       ("Выручка, грн", "revenue"), ("Доля выручки, %", "share"),
                                                       ("Прибыль, грн", "profit"), ("≈", "approx"), ("Отмен", "canceled")])],
    "store": [("Расходы ИИ", "ai", [("Модель", "model"), ("Действие", "action"), ("Запросов", "requests"),
                                    ("Токенов", "tokens"), ("$", "usd"), ("грн", "uah")])],
}


QUALITY_LABELS = {"no_photo": "Без фото", "no_description": "Без описания", "no_ua": "Без украинского названия",
                  "errors": "С ошибками заполнения"}


def xlsx(report: dict) -> bytes:
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    for title, key, cols in XLSX_TABLES[report["tab"]]:
        ws = wb.create_sheet(title[:31])
        ws.append([f"{report['from']} — {report['to']}"])
        ws.append([c for c, _ in cols])
        for r in report.get(key) or []:
            ws.append(["≈" if k == "approx" and r.get(k) else "" if k == "approx" else
                       (r.get(k) if r.get(k) is not None else "") for _, k in cols])
    if report["tab"] == "store":
        ws = wb.create_sheet("Магазин")
        ws.append(["Товаров", report["total"]])
        for k, n in report["statuses"].items():
            ws.append([products.STATUSES.get(k, k), n])
        for k, n in report["quality"].items():
            ws.append([QUALITY_LABELS[k], n])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
